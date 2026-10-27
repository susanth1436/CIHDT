"""Cover, Declaration, Certificate, Acknowledgements and the front lists.

Every size, indent and spacing value here was read out of the supplied
template rather than chosen, so the typography matches the specification
exactly: the three front-matter headings are 14 pt bold centred with a thick
underline, declaration and certificate bodies are justified at line spacing
1.5 with a 0.33 in left indent and a 0.50 in first line, and acknowledgement
paragraphs are justified at 1.5 with a 0.26 in first line and no left indent.
"""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from docx_kit import (bookmark, cell_text, clear_cell_borders,
                      front_heading, para)

C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT
R = WD_ALIGN_PARAGRAPH.RIGHT
J = WD_ALIGN_PARAGRAPH.JUSTIFY

DEAN = "Dr. Jaisankar N"
HOD = "<< Head of the Department >>"       # supply before submission
PROGRAMME = "Computer Science and Engineering"
INSTITUTE = "Vellore Institute of Technology, Vellore - 632 014"


def _gap(doc, pts=12):
    para(doc, "", size=11, line=1.0, after=pts)


# ===========================================================================
# Cover
# ===========================================================================
def cover(doc, *, title, team, guide, guide_designation, school, month):
    para(doc, "B.Tech. BCSE497J - Project-I", size=18, bold=True, align=C,
         line=1.5, before=2.85, after=22)
    para(doc, title.upper(), size=18, bold=True, align=C, line=1.5,
         before=2.85, after=24)

    para(doc, "Submitted in partial fulfillment of the requirements for the "
              "degree of", size=14, italic=True, align=C, line=1.5, after=10)
    para(doc, "Bachelor of Technology", size=19, bold=True, align=C,
         line=1.5, after=4)
    para(doc, "in", size=19, italic=True, align=C, line=1.5, after=4)
    para(doc, PROGRAMME, size=19, bold=True, align=C, line=1.5, after=18)

    para(doc, "by", size=14, italic=True, align=C, line=1.5, after=10)

    # Register number beside name, as the template's two-column block shows.
    table = doc.add_table(rows=len(team), cols=2)
    table.alignment = 1                                    # centre the table
    for r, (reg, name) in enumerate(team):
        for c, (text, align, bold) in enumerate(
                ((reg, R, True), (name.upper(), L, True))):
            cell = table.cell(r, c)
            cell.width = Inches(1.6 if c == 0 else 2.9)
            clear_cell_borders(cell)
            cell_text(cell, text, size=16, bold=bold, align=align, line=1.5)
    _gap(doc, 14)

    para(doc, "Under the Supervision of", size=14, italic=True, align=C,
         line=1.5, after=10)
    para(doc, guide, size=16, bold=True, align=C, line=1.5, after=2)
    para(doc, guide_designation, size=13, align=C, line=1.5, after=2)
    para(doc, school, size=13, align=C, line=1.5, after=20)

    para(doc, INSTITUTE, size=13, align=C, line=1.5, after=18)
    para(doc, month, size=13, align=C, line=1.5, after=0)


# ===========================================================================
# Declaration
# ===========================================================================
def _front_body(doc, text, *, before=4.5):
    """The declaration / certificate paragraph shape from the template."""
    return para(doc, text, size=12, align=J, line=1.5, before=before,
                after=0, left=0.33, indent=0.50)


def declaration(doc, *, title, team, guide, school):
    front_heading(doc, "DECLARATION", bookmark_name="bm_declaration")

    names = ", ".join(f"{name.upper()} ({reg})" for reg, name in team)
    _front_body(doc,
        f"We hereby declare that the project entitled “{title}” "
        f"submitted by us, for the award of the degree of Bachelor of "
        f"Technology in {PROGRAMME} to Vellore Institute of Technology, "
        f"Vellore, is a record of bonafide work carried out by us under the "
        f"supervision of {guide}, {school}.")

    _front_body(doc,
        "We further declare that the work reported in this project has not "
        "been submitted and will not be submitted, either in part or in "
        "full, for the award of any other degree or diploma in this "
        "institute or any other institute or university.", before=14)

    _gap(doc, 34)
    para(doc, "Place\t: Vellore", size=12, align=L, line=1.5, before=4.5,
         after=0, left=0.33)
    para(doc, "Date\t:", size=12, align=L, line=1.5, before=0, after=0,
         left=0.33)

    _gap(doc, 26)
    para(doc, "Signature of the Candidates", size=12, bold=True, align=L,
         line=1.5, before=0.1, after=6, left=4.18)
    for reg, name in team:
        para(doc, f"{name.upper()}  ({reg})", size=12, align=L, line=1.5,
             before=0, after=2, left=4.18)


# ===========================================================================
# Certificate
# ===========================================================================
def certificate(doc, *, title, team, guide, guide_designation, school):
    front_heading(doc, "CERTIFICATE", bookmark_name="bm_certificate")

    names = ", ".join(f"{name.upper()} ({reg})" for reg, name in team)
    _front_body(doc,
        f"This is to certify that the project entitled “{title}” "
        f"submitted by {names}, {school}, Vellore Institute of Technology, "
        f"Vellore, for the award of the degree of Bachelor of Technology in "
        f"{PROGRAMME}, is a record of bonafide work carried out by them "
        f"under my supervision, as per the VIT code of academic and research "
        f"ethics.")

    _front_body(doc,
        "The contents of this report have not been submitted and will not be "
        "submitted either in part or in full, for the award of any other "
        "degree or diploma in this institute or any other institute or "
        "university. The project fulfills the requirements and regulations "
        "of the University and in my opinion meets the necessary standards "
        "for submission.", before=14)

    _gap(doc, 30)
    para(doc, "Place\t: Vellore", size=12, align=L, line=1.5, before=8.05,
         after=0, left=0.33)
    para(doc, "Date\t:", size=12, align=L, line=1.5, before=6.85, after=0,
         left=0.33)

    _gap(doc, 22)
    para(doc, "Signature of the Guide", size=12, bold=True, align=L,
         line=1.5, before=6.85, after=2, left=0.33)
    para(doc, guide, size=12, align=L, line=1.5, before=0, after=0, left=0.33)
    para(doc, guide_designation, size=12, align=L, line=1.5, before=0,
         after=0, left=0.33)

    _gap(doc, 30)
    para(doc, "Examiner(s)", size=12, bold=True, align=R, line=1.5, after=26)

    para(doc, HOD, size=12, bold=True, align=C, line=1.5, after=0)
    para(doc, PROGRAMME, size=12, bold=True, align=C, line=1.5, after=0)


# ===========================================================================
# Acknowledgements
# ===========================================================================
def acknowledgements(doc, *, guide, team):
    front_heading(doc, "ACKNOWLEDGEMENTS",
                  bookmark_name="bm_acknowledgements")

    def ack(text):
        return para(doc, text, size=12, align=J, line=1.5, before=0,
                    after=10, indent=0.26)

    ack("We are deeply grateful to the management of Vellore Institute of "
        "Technology (VIT) for providing us with the opportunity and the "
        "resources to undertake this project. The laboratory facilities and "
        "the academic environment made it possible to build and evaluate the "
        "system described in this report.")

    ack(f"Our sincere thanks to {DEAN}, the Dean of the School of Computer "
        f"Science and Engineering (SCOPE), for his unwavering support and "
        f"encouragement throughout the course of this work.")

    ack(f"We express our profound appreciation to {HOD}, the Head of the "
        f"Department, for the insightful guidance and continuous "
        f"encouragement that shaped the direction of this project.")

    ack(f"We are immensely thankful to our project guide, {guide}, for his "
        f"dedicated mentorship and invaluable feedback. His patience, "
        f"knowledge and willingness to question our assumptions improved "
        f"this work at every stage, and several of the design decisions "
        f"defended in this report exist because of that scrutiny.")

    ack("We also thank the faculty members of SCOPE for their constructive "
        "comments during the review presentations, and our families and "
        "friends for their steady encouragement.")

    ack("Thank you all for your contributions and support.")

    _gap(doc, 26)
    para(doc, "Name of the Candidates", size=12, bold=True, align=L,
         line=1.5, before=0.1, after=6, left=4.18)
    for reg, name in team:
        para(doc, f"{name.upper()}  ({reg})", size=12, align=L, line=1.5,
             before=0, after=2, left=4.18)


# ===========================================================================
# Front lists
# ===========================================================================
def list_of_figures(doc, figures):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "List of Figures", size=16, bold=True, align=C, line=1.5,
         after=14)
    _simple_table(doc, ("Figure No.", "Title", "Page No."), figures)
    para(doc, "In the chapters, a figure caption appears below the figure "
              "and a table caption above the table. Captions are Times New "
              "Roman 10, centred.",
         size=10, italic=True, align=C, line=1.0, before=10, after=0)


def list_of_tables(doc, tables):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "List of Tables", size=16, bold=True, align=C, line=1.5,
         after=14)
    _simple_table(doc, ("Table No.", "Title", "Page No."), tables)


def _simple_table(doc, head, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    widths = (1.05, 4.35, 1.05)
    for c, text in enumerate(head):
        cell = table.cell(0, c)
        cell.width = Inches(widths[c])
        clear_cell_borders(cell)
        cell_text(cell, text, size=12, bold=True, align=C)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Inches(widths[c])
            clear_cell_borders(cell)
            cell_text(cell, value, size=12,
                      align=C if c != 1 else L)


def symbols_and_notations(doc, entries):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "Symbols and Notations", size=16, bold=True, align=C,
         line=1.5, after=14)
    for symbol, meaning in entries:
        p = para(doc, "", size=12, align=L, line=1.5, before=0, after=4,
                 left=0.36)
        from docx_kit import _style_run
        _style_run(p.add_run(f"{symbol}	"), size=12, bold=True, italic=True)
        _style_run(p.add_run(meaning), size=12)


def list_of_abbreviations(doc, entries):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "List of Abbreviations", size=16, bold=True, align=C,
         line=1.5, after=14)
    for short, long in entries:
        p = para(doc, "", size=12, align=L, line=1.5, before=0, after=4,
                 left=0.33)
        from docx_kit import _style_run
        _style_run(p.add_run(f"{short}\t"), size=12, bold=True)
        _style_run(p.add_run(long), size=12)
    para(doc, "Abbreviations are listed in alphabetical order.", size=10,
         italic=True, align=C, line=1.0, before=12, after=0)
